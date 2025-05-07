from typing_extensions import Annotated
import os
from mcp.server.fastmcp import FastMCP
from qdrant_client import QdrantClient
from qdrant_client.http.models import Filter, FieldCondition, MatchValue
from uuid import uuid4
from qdrant_client.models import PointStruct, VectorParams, Distance, PayloadSchemaType
import pdfplumber
from docx import Document as DocxDocument
from langchain.text_splitter import RecursiveCharacterTextSplitter
import traceback
import datetime
import tempfile
import json
import threading
import time
import speech_recognition as sr

# Create MCP instance
mcp = FastMCP("Qdrant_MCP")

# Placeholder embedding function - replace with proper embedding model later
def encode_text(text):
    return [0.0] * 768  # placeholder vector of size 768

@mcp.tool()
def create_qdrant_collection(
    collection_name: Annotated[str, "Name of the Qdrant collection to create"]
) -> str:
    """Creates a new collection in Qdrant with default vector settings."""

    try:
        # Initialize Qdrant client
        qdrant_client = QdrantClient(
            url=os.getenv("QDRANT_URL"),
            api_key=os.getenv("QDRANT_API_KEY")
        )

        # Check if collection already exists
        existing = qdrant_client.get_collections().collections
        if any(col.name == collection_name for col in existing):
            return f"Collection '{collection_name}' already exists."

        # First create collection with standard settings
        qdrant_client.create_collection(
            collection_name=collection_name,
            vectors_config=VectorParams(size=768, distance=Distance.COSINE)
        )
        
        # Then create the payload index separately
        try:
            qdrant_client.create_payload_index(
                collection_name=collection_name,
                field_name="document_name",
                field_schema="keyword"
            )
        except Exception as index_error:
            return f"Collection created but index creation failed: {str(index_error)}"
            
        return f"Collection '{collection_name}' created successfully with document_name index."

    except Exception as e:
        return f"Failed to create collection: {str(e)}"

@mcp.tool()
def create_payload_index(
    collection_name: Annotated[str, "Name of the collection to update"],
    field_name: Annotated[str, "Name of the field to index"]
) -> str:
    """Creates a payload index on the specified field."""
    try:
        qdrant_client = QdrantClient(
            url=os.getenv("QDRANT_URL"),
            api_key=os.getenv("QDRANT_API_KEY")
        )
        
        qdrant_client.create_payload_index(
            collection_name=collection_name,
            field_name=field_name,
            field_schema="keyword"
        )
        
        return f"Successfully created index on field '{field_name}' in collection '{collection_name}'."
    except Exception as e:
        return f"Failed to create index: {str(e)}"

@mcp.tool()
def split_file_into_chunks(
    file_path: Annotated[str, "Path to the input document (PDF, DOCX, or TXT)"]
) -> list[str]:
    """Splits a document into text chunks for embedding."""

    if not os.path.exists(file_path):
        return ["File not found."]

    ext = os.path.splitext(file_path)[1].lower()

    try:
        # Extract text
        if ext == ".pdf":
            with pdfplumber.open(file_path) as pdf:
                text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        elif ext == ".docx":
            doc = DocxDocument(file_path)
            text = "\n".join(p.text for p in doc.paragraphs)
        elif ext == ".txt":
            with open(file_path, "r", encoding="utf-8") as f:
                text = f.read()
        else:
            return ["Unsupported file type. Use .pdf, .docx, or .txt."]

        # Split into chunks
        splitter = RecursiveCharacterTextSplitter(chunk_size=500, chunk_overlap=50)
        chunks = splitter.split_text(text)

        return chunks if chunks else ["No content found in document."]
    except Exception as e:
        return [f"Error reading or splitting document: {str(e)}"]

@mcp.tool()
def upsert_chunks_to_qdrant(
    chunks: Annotated[list[str], "List of text chunks to upsert"],
    collection_name: Annotated[str, "Qdrant collection name"],
    document_name: Annotated[str, "Original document name (e.g., 'pay.docx')"]
) -> str:
    """Upserts text chunks into a Qdrant collection, tagging each chunk with the document name."""

    try:
        qdrant_client = QdrantClient(
            url=os.getenv("QDRANT_URL"),
            api_key=os.getenv("QDRANT_API_KEY")
        )

        # list of PointStructs with payloads
        points = []
        for chunk in chunks:
            vector = encode_text(chunk) 
            points.append(PointStruct(
                id=str(uuid4()),
                vector=vector,
                payload={
                    "text": chunk,
                    "document_name": document_name
                }
            ))

        # Upsert all points to the specified collection
        qdrant_client.upsert(
            collection_name=collection_name,
            points=points
        )

        return f"Upserted {len(points)} chunks from '{document_name}' into collection '{collection_name}'."

    except Exception as e:
        return f"Upsert failed: {str(e)}"

@mcp.tool()
def list_documents_in_collection(
    collection_name: Annotated[str, "Name of the Qdrant collection"]
) -> list:
    """Lists all unique document names in a collection."""
    
    try:
        qdrant_client = QdrantClient(
            url=os.getenv("QDRANT_URL"),
            api_key=os.getenv("QDRANT_API_KEY")
        )
        
        # Scroll through collection to find all document names
        results = qdrant_client.scroll(
            collection_name=collection_name,
            limit=10000,
            with_payload=["document_name"],
            with_vectors=False
        )[0]
        
        # Gives unique document names
        document_names = set()
        for point in results:
            doc_name = point.payload.get("document_name")
            if doc_name:
                document_names.add(doc_name)
        
        return sorted(list(document_names))
    
    except Exception as e:
        # Print more detailed error for debugging
        print(f"Error in list_documents_in_collection: {str(e)}")
        print(traceback.format_exc())
        return [f"Error listing documents: {str(e)}"]

@mcp.tool()
def retrieve_document_content(
    document_name: Annotated[str, "Document name to retrieve"],
    collection_name: Annotated[str, "Name of the Qdrant collection"]
) -> str:
    """Retrieves all text chunks from a specific document and combines them."""
    
    try:
        # Initialize Qdrant client
        qdrant_client = QdrantClient(
            url=os.getenv("QDRANT_URL"),
            api_key=os.getenv("QDRANT_API_KEY")
        )
        
        # Define filter for the specific document
        document_filter = Filter(
            must=[
                FieldCondition(
                    key="document_name",
                    match=MatchValue(value=document_name)
                )
            ]
        )
        
        # Scroll through all points with this document name
        results = qdrant_client.scroll(
            collection_name=collection_name,
            scroll_filter=document_filter,
            limit=10000,  # High limit to get all chunks
            with_payload=True,  # Get the full payload including text
            with_vectors=False  # No need for vectors
        )[0]  # First element is points, second is next_page_offset
        
        if not results:
            return f"No content found for document '{document_name}' in collection '{collection_name}'."
        
        # Extract text from each chunk and join them
        chunks = []
        for point in results:
            if point.payload and "text" in point.payload:
                chunks.append(point.payload["text"])
        
        # Join all chunks with double newlines
        full_content = "\n\n".join(chunks)
        
        return full_content if full_content else f"Document '{document_name}' exists but contains no text."
    
    except Exception as e:
        # Print more detailed error for debugging
        print(f"Error in retrieve_document_content: {str(e)}")
        print(traceback.format_exc())
        return f"Error retrieving document content: {str(e)}"

@mcp.tool()
def semantic_search(
    query: Annotated[str, "Search query text"],
    collection_name: Annotated[str, "Name of the Qdrant collection to search"],
    limit: Annotated[int, "Maximum number of results to return"] = 5
) -> list:
    """Searches for semantically similar content using the query text."""
    
    try:
        # Initialize Qdrant client
        qdrant_client = QdrantClient(
            url=os.getenv("QDRANT_URL"),
            api_key=os.getenv("QDRANT_API_KEY")
        )
        
        query_vector = encode_text(query)
        
        results = qdrant_client.search(
            collection_name=collection_name,
            query_vector=query_vector,
            limit=limit,
            with_payload=True
        )
        
        # Format results
        formatted_results = []
        for point in results:
            formatted_results.append({
                "id": str(point.id),
                "score": round(point.score, 4),
                "text": point.payload.get("text", ""),
                "document_name": point.payload.get("document_name", "unknown")
            })
        
        return formatted_results
    
    except Exception as e:
        print(f"Error in semantic_search: {str(e)}")
        print(traceback.format_exc())
        return [{"error": f"Search failed: {str(e)}"}]

@mcp.tool()
def delete_document(
    document_name: Annotated[str, "Name of document to delete"],
    collection_name: Annotated[str, "Name of the Qdrant collection"]
) -> str:
    """Deletes all chunks belonging to a specific document."""
    
    try:
        # Initialize Qdrant client
        qdrant_client = QdrantClient(
            url=os.getenv("QDRANT_URL"),
            api_key=os.getenv("QDRANT_API_KEY")
        )
        
        # Define filter for the specific document
        document_filter = Filter(
            must=[
                FieldCondition(
                    key="document_name",
                    match=MatchValue(value=document_name)
                )
            ]
        )
        
        # Delete points with this document name
        operation_result = qdrant_client.delete(
            collection_name=collection_name,
            points_selector=document_filter
        )
        
        return f"Successfully deleted document '{document_name}' from collection '{collection_name}'."
    
    except Exception as e:
        return f"Error deleting document: {str(e)}"

@mcp.tool()
def delete_qdrant_collection(
    collection_name: Annotated[str, "Name of the Qdrant collection to delete"]
) -> str:
    """Deletes an entire collection from Qdrant."""
    
    try:
        qdrant_client = QdrantClient(
            url=os.getenv("QDRANT_URL"),
            api_key=os.getenv("QDRANT_API_KEY")
        )
        
        # Check if collection exists before attempting to delete
        existing = qdrant_client.get_collections().collections
        if not any(col.name == collection_name for col in existing):
            return f"Collection '{collection_name}' does not exist."
        
        # Delete the collection
        qdrant_client.delete_collection(collection_name=collection_name)
        
        return f"Collection '{collection_name}' has been successfully deleted."
    
    except Exception as e:
        return f"Failed to delete collection: {str(e)}"
    
@mcp.tool()
def list_collections_with_documents() -> dict:
    """Lists all collections in the Qdrant database along with documents in each collection."""
    
    try:
        qdrant_client = QdrantClient(
            url=os.getenv("QDRANT_URL"),
            api_key=os.getenv("QDRANT_API_KEY")
        )
        
        # Get all collections
        collections_info = qdrant_client.get_collections()
        collection_names = [col.name for col in collections_info.collections]
        
        if not collection_names:
            return {"status": "No collections found."}
        
        # Create a dictionary to store collection 
        collections_with_documents = {}
        
        # For each collection gets the unique document names
        for collection_name in collection_names:
            try:
                # Scroll through collection to find all document names
                results = qdrant_client.scroll(
                    collection_name=collection_name,
                    limit=10000,
                    with_payload=["document_name"],
                    with_vectors=False
                )[0]
                
                # Extract unique document names
                document_names = set()
                for point in results:
                    doc_name = point.payload.get("document_name")
                    if doc_name:
                        document_names.add(doc_name)
                
                collections_with_documents[collection_name] = sorted(list(document_names))
                
            except Exception as collection_error:
                collections_with_documents[collection_name] = f"Error retrieving documents: {str(collection_error)}"
        
        return collections_with_documents
    
    except Exception as e:
        print(f"Error in list_collections_with_documents: {str(e)}")
        print(traceback.format_exc())
        return {"error": f"Error listing collections: {str(e)}"}



@mcp.tool()
def download_document(
    document_name: Annotated[str, "Name of the document to download"],
    collection_name: Annotated[str, "Name of the Qdrant collection"],
    output_format: Annotated[str, "Output format (txt, docx)"] = "txt",
    output_path: Annotated[str, "Custom output directory path (optional)"] = None
) -> str:
    """
    Extracts document content from Qdrant and saves it to the specified location.
    
    If output_path is provided, document will be saved to that location.
    Otherwise, it will use the default './downloads' directory.
    """
    try:
        # Initialize Qdrant client
        qdrant_client = QdrantClient(
            url=os.getenv("QDRANT_URL"),
            api_key=os.getenv("QDRANT_API_KEY")
        )
        
        # Define filter for the specific document
        document_filter = Filter(
            must=[
                FieldCondition(
                    key="document_name", 
                    match=MatchValue(value=document_name)
                )
            ]
        )
        
        results = qdrant_client.scroll(
            collection_name=collection_name,
            scroll_filter=document_filter,
            limit=10000,  
            with_payload=True,  
            with_vectors=False  
        )[0]
        
        if not results:
            return f"No content found for document '{document_name}' in collection '{collection_name}'."
        
        # Extract text from each chunk and join them
        chunks = []
        for point in results:
            if point.payload and "text" in point.payload:
                chunks.append(point.payload["text"])
        
        # Join all chunks with double newlines
        full_content = "\n\n".join(chunks)
        
        if not full_content:
            return f"Document '{document_name}' exists but contains no text."
        
        # Determine the output directory path
        if output_path:
            # Use the provided output path
            output_dir = output_path
        else:
            # Use the default download directory
            output_dir = "./downloads"
            os.makedirs(output_dir, exist_ok=True)
        
        # Ensure the output directory exists
        if not os.path.exists(output_dir):
            return f"Error: Output directory '{output_dir}' does not exist."
        
        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        base_name = os.path.splitext(document_name)[0]
        
        if output_format.lower() == "docx":
            file_path = os.path.join(output_dir, f"{base_name}_{timestamp}.docx")
            # Create a DOCX file
            doc = DocxDocument()
            # Split by double newlines to create paragraphs
            paragraphs = full_content.split("\n\n")
            for para in paragraphs:
                if para.strip():  # Skip empty paragraphs
                    doc.add_paragraph(para)
            doc.save(file_path)
            file_type = "DOCX"
        else:
            # Default to TXT
            file_path = os.path.join(output_dir, f"{base_name}_{timestamp}.txt")
            with open(file_path, "w", encoding="utf-8") as f:
                f.write(full_content)
            file_type = "TXT"
            
        return f"Document '{document_name}' has been exported to {file_type} format at: {file_path}"
        
    except Exception as e:
        print(f"Error in download_document: {str(e)}")
        print(traceback.format_exc())
        return f"Error exporting document: {str(e)}"

@mcp.tool()
def voice_to_doc_search(
    collection_name: Annotated[str, "Name of the Qdrant collection to search"],
    document_name: Annotated[str, "Name of the document to search within"],
    limit: Annotated[int, "Maximum number of results to return"] = 3
) -> dict:
    """
    Captures voice input and uses it to search within a specified document.
    Returns the most relevant text chunks from the document.
    """
    try:
        # Initialize speech recognizer
        recognizer = sr.Recognizer()
        
        # Capture audio from microphone
        with sr.Microphone() as source:
            print("Speak now.")
            # Adjust for ambient noise
            recognizer.adjust_for_ambient_noise(source, duration=1)
            # Record audio
            audio = recognizer.listen(source, timeout=5)
            print("Voice captured.")
        
        # Convert speech to text
        query = recognizer.recognize_google(audio)
        print(f"Search query: '{query}'")
        
        # Initialize Qdrant client
        qdrant_client = QdrantClient(
            url=os.getenv("QDRANT_URL"),
            api_key=os.getenv("QDRANT_API_KEY")
        )
        
        # Create query vector
        query_vector = encode_text(query)
        
        # Create filter for specific document
        document_filter = Filter(
            must=[
                FieldCondition(
                    key="document_name",
                    match=MatchValue(value=document_name)
                )
            ]
        )
        
        # Perform search
        results = qdrant_client.search(
            collection_name=collection_name,
            query_vector=query_vector,
            limit=limit,
            with_payload=True,
            query_filter=document_filter
        )
        
        # Format results
        search_results = []
        for point in results:
            search_results.append({
                "score": round(point.score, 4),
                "text": point.payload.get("text", ""),
                "document_name": point.payload.get("document_name", "unknown")
            })
        
        return {
            "voice_query": query,
            "document": document_name,
            "collection": collection_name,
            "results": search_results
        }
    
    except sr.RequestError as e:
        return {"error": f"Speech recognition service error: {str(e)}"}
    except sr.UnknownValueError:
        return {"error": "Unable to recognize speech"}
    except Exception as e:
        print(f"Error in voice_to_doc_search: {str(e)}")
        print(traceback.format_exc())
        return {"error": f"Search failed: {str(e)}"}

# Run the MCP server
if __name__ == "__main__":
    mcp.settings.port = 8201  
    mcp.settings.sse_path = "/qdrant_docx" 
    mcp.run(transport="sse")